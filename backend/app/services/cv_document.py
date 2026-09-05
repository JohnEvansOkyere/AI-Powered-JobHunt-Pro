"""Render an editable CV draft to an in-memory DOCX export."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from app.services.cv_tailoring import CVContent


def _add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(7, 86, 71)


def render_cv_docx(content: dict) -> bytes:
    cv = CVContent.model_validate(content)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name.add_run(cv.personal_info.name or "Curriculum Vitae")
    name_run.bold = True
    name_run.font.size = Pt(18)
    name_run.font.color.rgb = RGBColor(7, 86, 71)

    contact_values = [
        cv.personal_info.email,
        cv.personal_info.phone,
        cv.personal_info.location,
        cv.personal_info.linkedin,
        cv.personal_info.github,
        cv.personal_info.website,
    ]
    contact = document.add_paragraph("  |  ".join(value for value in contact_values if value))
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(8)

    if cv.summary:
        _add_heading(document, "Professional summary")
        document.add_paragraph(cv.summary)

    if cv.skills.technical or cv.skills.certifications or cv.skills.languages:
        _add_heading(document, "Skills")
        if cv.skills.technical:
            document.add_paragraph("Technical: " + ", ".join(cv.skills.technical))
        if cv.skills.certifications:
            document.add_paragraph("Certifications: " + ", ".join(cv.skills.certifications))
        if cv.skills.languages:
            document.add_paragraph("Languages: " + ", ".join(cv.skills.languages))

    if cv.experience:
        _add_heading(document, "Experience")
        for exp in cv.experience:
            header = document.add_paragraph()
            title = header.add_run(f"{exp.title} | {exp.company}".strip(" |"))
            title.bold = True
            dates = " - ".join(value for value in (exp.start_date, exp.end_date) if value)
            meta = " | ".join(value for value in (exp.location, dates) if value)
            if meta:
                header.add_run(f"\n{meta}").italic = True
            if exp.description:
                document.add_paragraph(exp.description)
            for achievement in exp.achievements:
                document.add_paragraph(achievement, style="List Bullet")

    if cv.education:
        _add_heading(document, "Education")
        for edu in cv.education:
            line = " | ".join(value for value in (edu.degree, edu.institution) if value)
            paragraph = document.add_paragraph()
            paragraph.add_run(line).bold = True
            meta = " | ".join(value for value in (edu.location, edu.graduation_date, edu.gpa) if value)
            if meta:
                paragraph.add_run(f"\n{meta}")

    if cv.projects:
        _add_heading(document, "Projects")
        for project in cv.projects:
            paragraph = document.add_paragraph()
            paragraph.add_run(project.name).bold = True
            if project.description:
                paragraph.add_run(f" — {project.description}")
            if project.technologies:
                document.add_paragraph(", ".join(project.technologies))

    output = BytesIO()
    document.save(output)
    return output.getvalue()
